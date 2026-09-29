import re, io as _io, sys

SRC = 'fast_server.py'
with open(SRC, encoding='utf-8') as f:
    code = f.read()
orig = len(code)

# ============================================================
# PATCH H : /upload/clear reset compliance_criteria
# ============================================================
OLD_H = (
    '    global rag_index, documents, kag_graph, uploaded_files, uploaded_raw_texts\r\n'
    '    documents          = []\r\n'
    '    uploaded_files     = []\r\n'
    '    uploaded_raw_texts = []\r\n'
    '    rag_index          = None\r\n'
    '    kag_graph      = {"nodes": {}, "relations": {}}'
)
NEW_H = (
    '    global rag_index, documents, kag_graph, uploaded_files, uploaded_raw_texts, compliance_criteria\r\n'
    '    documents           = []\r\n'
    '    uploaded_files      = []\r\n'
    '    uploaded_raw_texts  = []\r\n'
    '    compliance_criteria = []\r\n'
    '    rag_index           = None\r\n'
    '    kag_graph           = {"nodes": {}, "relations": {}}'
)
if OLD_H in code:
    code = code.replace(OLD_H, NEW_H, 1)
    print('H OK : /upload/clear updated (CRLF)')
else:
    # Try LF version
    OLD_H_LF = OLD_H.replace('\r\n', '\n')
    NEW_H_LF = NEW_H.replace('\r\n', '\n')
    if OLD_H_LF in code:
        code = code.replace(OLD_H_LF, NEW_H_LF, 1)
        print('H OK : /upload/clear updated (LF)')
    else:
        print('H FAIL - dumping search window:')
        idx = code.find('uploaded_raw_texts = []')
        print(repr(code[idx-100:idx+200]))
        sys.exit(1)

# ============================================================
# PATCH I : Nouveaux endpoints /compliance/*
# ============================================================
ANCHOR_I = '@app.get("/mlops/stats")'
assert ANCHOR_I in code, 'MLOPS NOT FOUND'

NEW_ENDPOINTS = r'''
# ── Compliance Matrix Endpoints ────────────────────────────────────────────────

@app.get("/compliance/list")
async def compliance_list(
    category: str = Query(None, description="Filtrer par categorie"),
    mandatory: str = Query(None, description="Yes ou No"),
    limit: int = Query(500, description="Nombre max de criteres")
):
    """Liste les criteres de conformite extraits. Filtrage par category et mandatory."""
    results = compliance_criteria
    if category:
        results = [c for c in results if c["category"].lower() == category.lower()]
    if mandatory:
        results = [c for c in results if c["mandatory"].lower() == mandatory.lower()]
    return {
        "total":            len(compliance_criteria),
        "filtered":         len(results),
        "criteria":         results[:limit],
        "categories":       list({c["category"] for c in compliance_criteria}),
        "mandatory_count":  sum(1 for c in compliance_criteria if c["mandatory"] == "Yes"),
        "optional_count":   sum(1 for c in compliance_criteria if c["mandatory"] == "No"),
    }


@app.get("/compliance/stats")
async def compliance_stats():
    """Statistiques sur les criteres extraits par categorie."""
    by_cat = {}
    for c in compliance_criteria:
        by_cat.setdefault(c["category"], {"total": 0, "mandatory": 0, "optional": 0})
        by_cat[c["category"]]["total"] += 1
        if c["mandatory"] == "Yes":
            by_cat[c["category"]]["mandatory"] += 1
        else:
            by_cat[c["category"]]["optional"] += 1
    return {
        "total_criteria":  len(compliance_criteria),
        "mandatory_total": sum(1 for c in compliance_criteria if c["mandatory"] == "Yes"),
        "optional_total":  sum(1 for c in compliance_criteria if c["mandatory"] == "No"),
        "by_category":     by_cat,
        "sources":         list({c["source"] for c in compliance_criteria}),
        "kag_nodes":       len(kag_graph["nodes"]),
        "kag_relations":   len(kag_graph["relations"])
    }


@app.get("/compliance/export")
async def compliance_export(fmt: str = Query("csv", description="Format: csv ou excel")):
    """Exporte la Compliance Matrix.
    CSV  : colonnes ID|Categorie|Critere|Obligatoire|Source|Chunk|Statut de Conformite
    Excel: meme colonnes avec mise en forme (vert=Yes, orange=No)
    """
    if not compliance_criteria:
        raise HTTPException(status_code=404,
            detail="Aucun critere extrait. Uploadez un document RFP d abord.")

    COLS    = ["id","category","criterion","mandatory","source","chunk_id","compliance_status"]
    HDRS_FR = ["ID","Categorie","Critere","Obligatoire","Source","Chunk","Statut de Conformite"]

    if fmt.lower() in ("excel", "xlsx"):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Compliance Matrix"

            # Header
            fill_hdr = PatternFill("solid", fgColor="1F3A5F")
            font_hdr = Font(bold=True, color="FFFFFF")
            for ci, h in enumerate(HDRS_FR, 1):
                cell = ws.cell(row=1, column=ci, value=h)
                cell.fill = fill_hdr; cell.font = font_hdr
                cell.alignment = Alignment(horizontal="center")

            # Rows
            fill_yes = PatternFill("solid", fgColor="C6EFCE")
            fill_no  = PatternFill("solid", fgColor="FFEB9C")
            font_yes = Font(color="276221", bold=True)
            font_no  = Font(color="9C6500")
            for ri, crit in enumerate(compliance_criteria, 2):
                for ci, key in enumerate(COLS, 1):
                    ws.cell(row=ri, column=ci, value=crit.get(key, ""))
                mc = ws.cell(row=ri, column=4)
                if crit.get("mandatory") == "Yes":
                    mc.fill = fill_yes; mc.font = font_yes
                else:
                    mc.fill = fill_no;  mc.font = font_no

            for col, w in zip("ABCDEFG", [12, 20, 60, 14, 30, 26, 30]):
                ws.column_dimensions[col].width = w

            buf = _io.BytesIO()
            wb.save(buf); buf.seek(0)
            return StreamingResponse(buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition":"attachment; filename=compliance_matrix.xlsx"})
        except ImportError:
            pass   # fallback CSV

    # CSV (defaut, BOM pour Excel FR)
    buf = _io.StringIO()
    import csv as _csv
    writer = _csv.DictWriter(buf, fieldnames=COLS, extrasaction="ignore")
    writer.writerow(dict(zip(COLS, HDRS_FR)))
    for crit in compliance_criteria:
        writer.writerow(crit)
    buf.seek(0)
    return StreamingResponse(
        _io.BytesIO(buf.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition":"attachment; filename=compliance_matrix.csv"})


'''
code = code.replace(ANCHOR_I, NEW_ENDPOINTS + ANCHOR_I, 1)
print('I OK : /compliance/list, /stats, /export endpoints')

# ============================================================
# Validation syntaxe Python
# ============================================================
import py_compile, tempfile, os
tmp = SRC + '.tmp_validate.py'
with open(tmp, 'w', encoding='utf-8') as f:
    f.write(code)
try:
    py_compile.compile(tmp, doraise=True)
    os.remove(tmp)
    print('SYNTAX OK')
except py_compile.PyCompileError as e:
    os.remove(tmp)
    print(f'SYNTAX ERROR: {e}')
    sys.exit(1)

# ============================================================
# Ecriture finale
# ============================================================
with open(SRC, 'w', encoding='utf-8') as f:
    f.write(code)
print(f'DONE: {orig} -> {len(code)} bytes (+{len(code)-orig})')
